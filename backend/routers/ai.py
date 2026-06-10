from fastapi import APIRouter, HTTPException, Depends, status, UploadFile, File
import os
import json
import io
import pypdf
import docx
from pydantic import BaseModel
from typing import List, Optional
try:
    from backend.auth import get_current_user
    from backend.models import User
except ImportError:
    from auth import get_current_user
    from models import User

import google.generativeai as genai

router = APIRouter(prefix="/api/ai", tags=["AI"])

class AISuggestionRequest(BaseModel):
    class_name: str
    subject: str
    topic: str
    question_type: str
    count: int
    board: Optional[str] = "CBSE"

@router.post("/suggest-questions")
def suggest_questions(
    req: AISuggestionRequest,
    current_user: User = Depends(get_current_user)
):
    import datetime
    # Enforce subscription feature check
    features = []
    if current_user.subscription_plan and (current_user.subscription_expires_at is None or current_user.subscription_expires_at > datetime.datetime.utcnow()):
        features = current_user.subscription_plan.features or []
        
    if "ai_suggestions" not in features:
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail="AI Question Suggestion is not enabled on your subscription plan. Please upgrade to a Professional (Medium) or Enterprise (All Features) plan to access this feature."
        )

    from dotenv import load_dotenv
    env_path = os.path.join(os.path.dirname(os.path.dirname(__file__)), ".env")
    load_dotenv(dotenv_path=env_path)
    
    api_key = os.environ.get("GEMINI_API_KEY")
    if not api_key:
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f"Gemini API Key is not configured on the server. Tried loading from {env_path}"
        )

    genai.configure(api_key=api_key)
    
    # Construct specialized instructions based on question_type:
    type_instructions = ""
    
    if req.question_type == 'Paragraph and questions based on it':
        type_instructions = f"""
- This is a Comprehension/Passage reading question type.
- The 'question_text' must contain a reading passage/paragraph suitable for the subject and class (written in the correct script/language).
- The 'sub_questions' must contain exactly {req.count} questions based on that passage.
- Set 'marks' on the main question to the sum of sub-questions' marks.
- Set 'answer_text' to a brief overview of the passage or leave it empty.
- Each sub-question in the 'sub_questions' list must be structured as a dictionary:
  {{
      "text": "The sub-question text (e.g. '1. What did the author say about...')",
      "answer": "The answer/key points for this sub-question",
      "type": "text",
      "marks": 1,
      "options": []
  }}
"""
    elif req.question_type == 'Letter writing question':
        type_instructions = f"""
- This is a Letter Writing question.
- Generate a formal or informal letter writing scenario/prompt based on latest CBSE formats for Class {req.class_name}.
- Place the prompt in 'question_text'.
- Place the model format, marks division (e.g., Format: 1 mark, Content: 2 marks, Expression: 2 marks), and sample letter/key points in 'answer_text'.
- Ensure 'sub_questions' is an empty list [].
"""
    elif req.question_type == 'Essay':
        type_instructions = f"""
- This is an Essay (निबंध) writing question.
- Place the essay topic along with structural outline hints/points in 'question_text'.
- Place value points or key headers/sample content in 'answer_text'.
- Ensure 'sub_questions' is an empty list [].
"""
    elif req.question_type == 'Paragraph writing':
        type_instructions = f"""
- This is a Paragraph Writing prompt.
- Place the topic and key guidance points (usually 80-120 words) in 'question_text'.
- Place key points or model answers in 'answer_text'.
- Ensure 'sub_questions' is an empty list [].
"""
    elif req.question_type == 'Email':
        type_instructions = f"""
- This is an Email writing task.
- Place the email scenario/guidelines in 'question_text'.
- Place the correct email layout template and sample answer in 'answer_text'.
- Ensure 'sub_questions' is an empty list [].
"""
    elif req.question_type == 'Advertisement writing':
        type_instructions = f"""
- This is an Advertisement Writing (विज्ञापन लेखन) prompt.
- Place the prompt/scenario in 'question_text'.
- Place the model points (e.g. border frame, catchy title, contact, main details) in 'answer_text'.
- Ensure 'sub_questions' is an empty list [].
"""
    elif req.question_type == 'Notice writing':
        type_instructions = f"""
- This is a Notice Writing (सूचना लेखन) task.
- Place the notice scenario (school/org name, heading, date, content, signature block) in 'question_text'.
- Place the notice layout structure and answer key in 'answer_text'.
- Ensure 'sub_questions' is an empty list [].
"""
    elif req.question_type == 'Story writing':
        type_instructions = f"""
- This is a Story Writing (लघुकथा लेखन) prompt.
- Place outline hints, title prompt, or an opening sentence in 'question_text'.
- Place key plot ideas, moral, or outline details in 'answer_text'.
- Ensure 'sub_questions' is an empty list [].
"""
    elif req.question_type.startswith('Grammar:'):
        grammar_type = req.question_type.replace('Grammar:', '').strip()
        type_instructions = f"""
- This is a Grammar question focusing on '{grammar_type}'.
- Generate questions according to CBSE board syllabus for Class {req.class_name}.
- In 'question_text', provide the question clearly (e.g. 'निम्नलिखित शब्द का विलोम शब्द लिखिए:' or 'दिए गए मुहावरे का अर्थ लिखकर वाक्य में प्रयोग कीजिए:', etc.).
- In 'answer_text', provide the exact grammatical answer or dissolution.
- If it's a multiple choice grammar question (common in CBSE), you can set 'sub_questions' to represent the options (exactly 4 options as sub_questions) and the correct answer in 'answer_text'. Otherwise, leave 'sub_questions' as an empty list [].
"""
    elif 'MCQ' in req.question_type:
        type_instructions = """
- This is a Multiple Choice Question (MCQ).
- The 'question_text' must contain the main question stem.
- The 'sub_questions' list must contain EXACTLY 4 options, each as a dictionary:
  {"text": "Option text", "type": "MCQ", "options": []}
- Set 'answer_text' to the correct option text (must match one of the options).
"""
    else:
        type_instructions = f"""
- Generate questions of type '{req.question_type}'.
- Place the question text in 'question_text'.
- Place the answer or key solution points in 'answer_text'.
- Leave 'sub_questions' as an empty list [].
"""

    model = genai.GenerativeModel(
        'gemini-2.5-flash',
        generation_config={"response_mime_type": "application/json"}
    )
    
    board_name = req.board or "CBSE"
    board_guidelines = ""
    if board_name == "GSEB":
        board_guidelines = f"""
1. **Syllabus & Standard**: Align questions strictly with the Gujarat Secondary and Higher Secondary Education Board (GSEB) syllabus and pattern for Class {req.class_name}.
2. **GSEB SSC/HSC Style**: For Class 10 and Class 12, use GSEB board exam blueprints (e.g. Part A Objective / Part B Descriptive splits, with descriptive subdivided into 4 knowledge/understanding/application/skill sections).
3. **Language Mediums**: GSEB papers must use standard Gujarati Board vocabulary.
"""
    else:
        board_guidelines = f"""
1. **Syllabus & Standard**: Align questions strictly with the latest CBSE board curriculum and standard for Class {req.class_name} {req.subject}.
2. **Competency-Based**: Incorporate real-world scenarios, analytical thinking, or application-based inquiries where appropriate following latest CBSE guidelines.
"""

    prompt = f"""
You are an expert {board_name} teacher for Class {req.class_name} and Subject: {req.subject}.
Your task is to generate {req.count} questions of type '{req.question_type}' about the topic: '{req.topic}'.

### IMPORTANT {board_name} BOARD STYLE & CURRICULUM GUIDELINES:
{board_guidelines}
3. **Language & Script**: 
   - If the subject or topic is in Hindi or Sanskrit, write the content in the **Devanagari script** (हिन्दी/संस्कृत).
   - If the subject is Gujarati, write the content in the **Gujarati script** (ગુજરાતી).
   - Ensure perfect grammar, correct script spelling, and appropriate level of vocabulary.

### QUESTION TYPE SPECIFIC DIRECTIONS:
{type_instructions}

### JSON FORMAT REQUIREMENT:
Output ONLY a JSON array of objects. Do not include markdown blocks like ```json or any other text.
The JSON array should contain objects with the following structure:
{{
    "question_text": "The question text here (HTML formatting such as <b>, <i>, <sup>, <sub> is allowed)",
    "answer_text": "The correct answer, solution key, or outline",
    "marks": 1, // Appropriate marks for this question type
    "sub_questions": [] // Array of sub-question objects if specified above, otherwise empty
}}
"""

    try:
        response = model.generate_content(prompt)
        text_response = response.text.strip()
        
        # Cleanup in case the model returns markdown code blocks
        if text_response.startswith("```json"):
            text_response = text_response[7:]
        if text_response.endswith("```"):
            text_response = text_response[:-3]
            
        questions = json.loads(text_response.strip())
        return questions
    except Exception as e:
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f"Failed to generate questions: {str(e)}"
        )

@router.post("/import-paper")
def import_paper(
    file: UploadFile = File(...),
    current_user: User = Depends(get_current_user)
):
    import datetime
    # Enforce subscription feature check
    features = []
    if current_user.subscription_plan and (current_user.subscription_expires_at is None or current_user.subscription_expires_at > datetime.datetime.utcnow()):
        features = current_user.subscription_plan.features or []
        
    if "smart_scanner" not in features:
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail="Smart Paper Scanner (PDF & Word Import) is not enabled on your subscription plan. Please upgrade to the Enterprise (All Features) plan to access this feature."
        )

    import tempfile
    
    filename = file.filename or ""
    ext = os.path.splitext(filename)[1].lower()
    
    # We load Gemini API Key
    from dotenv import load_dotenv
    env_path = os.path.join(os.path.dirname(os.path.dirname(__file__)), ".env")
    load_dotenv(dotenv_path=env_path)
    
    api_key = os.environ.get("GEMINI_API_KEY")
    if not api_key:
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail="Gemini API Key is not configured on the server."
        )

    genai.configure(api_key=api_key)
    
    model = genai.GenerativeModel(
        'gemini-2.5-flash',
        generation_config={"response_mime_type": "application/json"}
    )
    
    prompt = """
You are an expert exam paper structure analyzer.
Your task is to extract all details from the provided exam paper and structure it into a precise, valid JSON object matching our database format.

### CRITICAL ACCURACY INSTRUCTIONS:
1. **NO OMISSIONS**: Extract every single question, passage, comprehension paragraph, option, and instruction. Never summarize, shorten, or omit any text.
2. **COMPREHENSION PASSAGES**: If a question contains a reading passage/comprehension paragraph (very common in language exams like Gujarati, Hindi, English), you MUST include the entire passage text inside the main question's `question_text`.
3. **SUB-QUESTIONS**: If a question has sub-parts (e.g. 1), 2), 3) or a), b), c)), you MUST extract the full text of each sub-question. Do not just output numbers like "1)" or "a)".
4. **GUJARATI / HINDI / SANSKRIT SCRIPT**: Ensure perfect unicode reproduction of Indian scripts. Never lose characters or accents.
5. **MCQ OPTIONS**: Ensure every MCQ option is fully extracted into the `sub_questions` (as type: "MCQ") or `options` array.
6. **NO SHORTENING**: Ensure no question content is replaced with placeholders like "..." or "same as above". Extract everything in full.

Your output MUST be a JSON object with the following exact keys:
1. "title": The title of the question paper (e.g., "Class X English Unit Test 1", "Quarterly Exam - Mathematics"). If not clearly defined, construct a descriptive title.
2. "subject": The subject name (e.g. "English", "Mathematics", "Science", "Gujarati", "Hindi").
3. "class_name": The standard or class name (e.g., "Class 10", "Class IX").
4. "max_marks": The maximum marks as an integer. Default to 25 if not found.
5. "time_duration": The time duration as a string. Optional/null if not found.
6. "date_str": The date of the exam as a string. Optional/null if not found.
7. "instructions": The general instructions for students, formatted in clean HTML. You may use simple lists (<ul>, <ol>, <li>) or paragraph tags (<p>).
8. "questions": An array of questions present in the paper. Order them exactly as they appear in the original text.

Each question object in the "questions" array MUST have:
- "section": The section or part header if present (e.g., "SECTION A", "PART B", "Choose the correct option:"). If the question falls under a section, fill this with that section header.
- "question_type": The type of question. Use standard type names like "MCQ", "Short Answer", "Long Answer", "True/False", "Letter Writing", etc.
- "question_text": The main text of the question (and the reading passage/comprehension text if applicable). HTML tags like <b>, <i>, <sup>, <sub>, and simple math symbols are allowed. If the question has a prefix like "Q1.", remove the "Q1." prefix as the renderer handles numbering automatically, but preserve the core question text.
- "answer_text": If the text contains the answer key or solutions, extract the correct option/answer text here. Otherwise, leave it as an empty string.
- "marks": The marks allocated to this question as an integer.
- "sub_questions": A list of sub-parts or options if it is an MCQ, or if the question contains sub-questions (e.g. Q1 has sub-questions (a), (b), (c)).
  Each sub-question object in this list MUST have:
  - "text": The sub-question text or the MCQ option text. If it is an MCQ option, just place the option content (e.g., "Paris").
  - "answer": The answer key/option solution if found.
  - "type": "MCQ" if it is an option of a multiple choice question, or "text" for regular sub-questions.
  - "options": For MCQ-type sub-questions, this is not needed or can be an empty array. If a sub-question itself is an MCQ, specify its options as an array of 4 strings. Otherwise, provide an empty array [].
  - "marks": Marks allocated to this sub-question.

IMPORTANT instructions for Gujarati/Hindi content:
- If the text is in Hindi/Sanskrit, preserve the Devanagari script.
- If the text is in Gujarati, preserve the Gujarati script.

Return ONLY the valid JSON object. Do not include markdown wraps like ```json.
"""

    file_bytes = file.file.read()
    
    if ext == ".pdf":
        local_text = ""
        try:
            pdf_file = io.BytesIO(file_bytes)
            reader = pypdf.PdfReader(pdf_file)
            text_pages = []
            for page in reader.pages:
                text_pages.append(page.extract_text() or "")
            local_text = "\n".join(text_pages).strip()
        except Exception:
            pass
            
        if len(local_text) > 120:
            try:
                # Send raw text to Gemini (takes only 3-5 seconds!)
                response = model.generate_content([
                    f"Here is the text extracted from the PDF:\n\n{local_text}\n\n{prompt}"
                ])
                text_response = response.text.strip()
            except Exception as e:
                raise HTTPException(
                    status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
                    detail=f"Failed to analyze PDF text with Gemini: {str(e)}"
                )
        else:
            try:
                # Fall back to visual processing (takes 15-25 seconds but handles images/scans)
                response = model.generate_content([
                    {
                        "mime_type": "application/pdf",
                        "data": file_bytes
                    },
                    prompt
                ])
                text_response = response.text.strip()
            except Exception as e:
                raise HTTPException(
                    status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
                    detail=f"Failed to visually analyze PDF with Gemini: {str(e)}"
                )
    elif ext in [".docx", ".doc"]:
        # DOCX: extract text locally and send to Gemini
        try:
            docx_file = io.BytesIO(file_bytes)
            doc = docx.Document(docx_file)
            text_parts = []
            for p in doc.paragraphs:
                text_parts.append(p.text)
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        text_parts.append(cell.text)
            content = "\n".join(text_parts)
        except Exception as e:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f"Failed to parse DOCX file: {str(e)}"
            )
            
        if not content.strip():
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Could not extract any text from the uploaded file."
            )
            
        try:
            response = model.generate_content([
                f"Here is the text extracted from the document:\n\n{content}\n\n{prompt}"
            ])
            text_response = response.text.strip()
        except Exception as e:
            raise HTTPException(
                status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
                detail=f"Failed to analyze DOCX with Gemini: {str(e)}"
            )
    else:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Unsupported file format. Please upload a PDF or DOCX file."
        )

    try:
        # Cleanup markdown formatting wraps if returned
        if text_response.startswith("```json"):
            text_response = text_response[7:]
        if text_response.endswith("```"):
            text_response = text_response[:-3]
            
        paper_data = json.loads(text_response.strip())
        return paper_data
    except Exception as e:
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f"Failed to parse Gemini JSON output: {str(e)}"
        )
